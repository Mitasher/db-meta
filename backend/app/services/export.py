import re
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from tempfile import SpooledTemporaryFile
from typing import IO, Any, Iterator

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from sqlalchemy import Engine
from sqlalchemy.exc import SQLAlchemyError

from app.db.models import MetaTable
from app.schemas.tables import DataQueryIn
from app.services.table_data import build_statements, source_error

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ExportTooLargeError(Exception):
    """Под фильтры попадает больше строк, чем разрешено выгружать."""


def excel_value(value: Any) -> Any:
    """Значение из БД в то, что openpyxl умеет записать в ячейку."""
    if value is None or isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, datetime) and value.tzinfo is not None:
        # Excel не знает часовых поясов
        return value.replace(tzinfo=None)
    if isinstance(value, (bytes, bytearray)):
        return "0x" + bytes(value).hex()
    if isinstance(value, str):
        # Управляющие символы openpyxl не пишет; 32767 — предел длины текста в ячейке Excel
        return ILLEGAL_CHARACTERS_RE.sub("", value)[:32767]
    if isinstance(value, (date, time, timedelta, Decimal)):
        return value
    return str(value)


def sheet_title(meta_table: MetaTable) -> str:
    """Имя листа: подпись таблицы без запрещённых символов, не длиннее 31 знака."""
    title = re.sub(r"[\[\]:*?/\\]", "_", meta_table.label or meta_table.table_name)
    return title[:31] or "data"


def export_filename(meta_table: MetaTable, created_at: datetime) -> str:
    """Имя файла: подпись таблицы и время выгрузки."""
    return f"{meta_table.label or meta_table.table_name}_{created_at:%Y-%m-%d_%H-%M}.xlsx"


def write_xlsx(engine: Engine, meta_table: MetaTable, query: DataQueryIn, max_rows: int) -> IO[bytes]:
    """Все строки по фильтрам запроса в xlsx; файл держится в памяти, крупный — уходит на диск."""
    statements = build_statements(meta_table, query)
    try:
        with engine.connect() as connection:
            total = connection.execute(statements.count).scalar_one()
            if total > max_rows:
                raise ExportTooLargeError(
                    f"Под фильтры попадает {total} строк, выгрузить можно не больше {max_rows} — сузьте фильтры"
                )
            workbook = Workbook(write_only=True)
            sheet = workbook.create_sheet(title=sheet_title(meta_table))
            sheet.append([column.label or column.column_name for column in statements.columns])
            # stream_results: строки идут из БД порциями, а не все разом в память
            result = connection.execution_options(stream_results=True).execute(statements.data.limit(max_rows))
            for row in result:
                sheet.append([excel_value(value) for value in row])
    except SQLAlchemyError as error:
        raise source_error(error) from error

    # До 20 МБ файл в памяти, дальше — во временном файле на диске
    output = SpooledTemporaryFile(max_size=20 * 1024 * 1024)
    workbook.save(output)
    output.seek(0)
    return output


def iter_file(output: IO[bytes]) -> Iterator[bytes]:
    """Отдаёт файл кусками и закрывает его в конце."""
    try:
        while chunk := output.read(64 * 1024):
            yield chunk
    finally:
        output.close()
